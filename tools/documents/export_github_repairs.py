#!/usr/bin/env python3
"""Export the repair records to DOCX/PDF (python-docx, reportlab, markdown, bs4).

Run with --font pointing to a Chinese TrueType font or TTC collection.
All input/output paths are derived from this repository; no source files change.
"""
import argparse
import html
import re
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
)

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'doc/work_records/github_repairs'
ORDER = ['README', 'workflow', 'triton_fp8_conversion', 'quack_autotuning_rng',
         'deepspeed_overflow', 'mooncake_queue_lifecycle',
         'cutlass_workspace_descriptor', 'vllm_padding_validation']
FLOW = '''
1. 固定失败条件和基线。
2. 在原版上复现目标失败。
3. 分析根因并实现最小补丁。
4. 用相同测试执行负正对照。
5. 提交可审阅差异及验证边界。
6. Review 若发现契约问题或回归，返回第 3 步；若需补证据，返回第 4 步。
7. 无新增修改时，等待维护者与 CI 流程。
'''


def target_url(value, source):
    if value.startswith(('https://', 'http://', '#')):
        return value
    path, _, anchor = value.partition('#')
    resolved = (source.parent / path).resolve()
    relative = resolved.relative_to(ROOT)
    return 'https://github.com/XFDG/ICS6201/blob/me/' + relative.as_posix() + (
        '#' + anchor if anchor else '')


def pdf_inline(node, source):
    if isinstance(node, NavigableString):
        return html.escape(str(node))
    content = ''.join(pdf_inline(child, source) for child in node.children)
    if node.name == 'br':
        return '<br/>'
    if node.name in ('strong', 'b'):
        return '<b>' + content + '</b>'
    if node.name in ('em', 'i'):
        return '<i>' + content + '</i>'
    if node.name == 'a':
        url = html.escape(target_url(node.get('href', ''), source), quote=True)
        return f'<a href="{url}" color="#12628a">{content}</a>'
    return content


def word_inline(paragraph, node, source, bold=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node))
        run.bold = bold
        return
    if node.name == 'br':
        paragraph.add_run().add_break()
        return
    if node.name == 'a':
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(
            target_url(node.get('href', ''), source), RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '12628A')
        props.append(color)
        run.append(props)
        text = OxmlElement('w:t')
        text.text = node.get_text()
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        return
    for child in node.children:
        word_inline(paragraph, child, source, bold or node.name in ('strong', 'b'))


def new_word():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.7)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3',
                 'List Bullet', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'WenQuanYi Zen Hei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'WenQuanYi Zen Hei')
        style.font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.2
    for name, size in [('Heading 1', 19), ('Heading 2', 14), ('Heading 3', 12)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('123B55')
    header = section.header.paragraphs[0]
    header.text = 'GitHub 修复开发记录 · 脱密版'
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    doc.core_properties.author = 'Work Records'
    return doc


def export(sources, output):
    doc = new_word()
    body = ParagraphStyle('body', fontName='CJK', fontSize=10, leading=15,
                          spaceAfter=7, wordWrap='CJK', alignment=TA_LEFT)
    heading = {n: ParagraphStyle('h'+str(n), parent=body, fontSize=size,
                leading=size*1.35, textColor=colors.HexColor('#123B55'),
                spaceBefore=10, spaceAfter=8, keepWithNext=True)
               for n, size in [(1, 18), (2, 13), (3, 11)]}
    cell_style = ParagraphStyle('cell', parent=body, fontSize=8.4, leading=12,
                                spaceAfter=0)
    code_style = ParagraphStyle('code', parent=body, fontSize=9, leading=13,
                                backColor=colors.HexColor('#F1F5F8'), borderPadding=6)
    story = []
    for index, source in enumerate(sources):
        if index:
            story.append(PageBreak())
            doc.add_page_break()
        text = source.read_text(encoding='utf-8')
        text = re.sub(r'```mermaid\n.*?```', FLOW, text, flags=re.S)
        soup = BeautifulSoup(markdown.markdown(text, extensions=['tables', 'fenced_code']),
                             'html.parser')
        for element in soup.children:
            if isinstance(element, NavigableString):
                continue
            tag = element.name
            if re.fullmatch('h[1-6]', tag):
                level = min(int(tag[1]), 3)
                doc.add_heading(element.get_text(), level=level)
                story.append(Paragraph(pdf_inline(element, source), heading[level]))
            elif tag == 'table':
                rows = [row.find_all(['th', 'td'], recursive=False)
                        for row in element.find_all('tr')]
                count = max(map(len, rows))
                table = doc.add_table(rows=0, cols=count)
                table.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    cells = table.add_row().cells
                    for ci, cell in enumerate(row):
                        word_inline(cells[ci].paragraphs[0], cell, source, ri == 0)
                    trpr = table.rows[-1]._tr.get_or_add_trPr()
                    trpr.append(OxmlElement('w:cantSplit'))
                    if ri == 0:
                        trpr.append(OxmlElement('w:tblHeader'))
                data = [[Paragraph(pdf_inline(cell, source), cell_style) for cell in row]
                        for row in rows]
                width = A4[0] - 96
                pdf_table = Table(data, colWidths=[width/count]*count, repeatRows=1,
                                  hAlign='LEFT')
                pdf_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E5EEF4')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#BCCCD7')),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ]))
                story.extend([pdf_table, Spacer(1, 9)])
            elif tag in ('ul', 'ol'):
                for li, item in enumerate(element.find_all('li', recursive=False), 1):
                    p = doc.add_paragraph(style='List Number' if tag == 'ol' else 'List Bullet')
                    word_inline(p, item, source)
                    prefix = str(li)+'. ' if tag == 'ol' else '• '
                    story.append(Paragraph(prefix+pdf_inline(item, source), body))
            elif tag == 'pre':
                code = element.get_text().strip('\n')
                p = doc.add_paragraph()
                p.add_run(code).font.size = Pt(9)
                story.append(Paragraph(html.escape(code).replace('\n', '<br/>'), code_style))
            else:
                p = doc.add_paragraph()
                word_inline(p, element, source)
                story.append(Paragraph(pdf_inline(element, source), body))
    doc.save(output.with_suffix('.docx'))

    def footer(canvas, pdf):
        canvas.saveState()
        canvas.setFont('CJK', 8)
        canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawString(48, A4[1]-28, 'GitHub 修复开发记录 · 脱密版')
        canvas.drawRightString(A4[0]-48, 25, f'{pdf.page}')
        canvas.restoreState()

    SimpleDocTemplate(str(output.with_suffix('.pdf')), pagesize=A4,
                      rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=42,
                      title='GitHub 修复开发记录 - '+output.stem,
                      author='Work Records').build(story, onFirstPage=footer, onLaterPages=footer)
    print(output.name, flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--font', type=Path,
                        default=Path('/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'))
    args = parser.parse_args()
    pdfmetrics.registerFont(TTFont('CJK', str(args.font), subfontIndex=0))
    pdfmetrics.registerFontFamily('CJK', normal='CJK', bold='CJK', italic='CJK', boldItalic='CJK')
    files = [SOURCE / (name+'.md') for name in ORDER]
    for source in files:
        export([source], source.with_suffix(''))
    export(files, SOURCE / 'GitHub修复开发记录_脱密合集')


if __name__ == '__main__':
    main()
