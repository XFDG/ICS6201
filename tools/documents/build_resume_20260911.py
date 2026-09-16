#!/usr/bin/env python3
"""Build the 2026-09-11 Chinese resume draft from the existing Word template."""

from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "doc" / "冯浩然_AiInfra香港中文大学_15024999885.docx"
OUTPUT = ROOT / "doc" / "冯浩然_AIInfra中文简历_RL训练推理算子版_20260911.docx"


RL_PARTS = [
    ("RL 算子支持：", True),
    (
        "面向 MoE-RL rollout 的训推一致性、异步解码与 CUDA Graph 卡死，负责 Router Replay 的 "
        "route/logprob 可观测及 A/B 指标闭环，实现 GPU token-history/Triton fused-hash 和 TP>1 "
        "batch-invariant 安全路径，并将 TP2 hang 定位到 FlashInfer Lamport sentinel 的 FTZ 误判、"
        "回移上游位级修复；20-step 对照中 replay target 达到 0 mismatch，fτ²/KL 分别降低 "
        "36–145×/4–7×；TP1 吞吐较 sync 提升 ",
        False,
    ),
    ("4.99%", True),
    ("，TP2/TP4 各 84 case 均 0 mismatch，decode 吞吐提升 ", False),
    ("4.6%/3.0%", True),
    ("，Graph capture/replay 稳定通过。", False),
]


TRAINING_PARTS = [
    ("训练算子支持：", True),
    (
        "针对长序列变长 SWA deterministic backward 异常，结合 Nsys/NCU 因子隔离确认 dQ semaphore "
        "依赖链占确定性增量 99.67%；在 FA3 原 fused main kernel 内对齐 reverse scheduler 与 "
        "contributor-relative ticket，并基于 host 已有长度元数据动态分流，不新增 kernel/workspace/D2H；"
        "代表 packed shape 完整 backward 由 6.755 ms 降至 1.699 ms（",
        False,
    ),
    ("3.98×", True),
    ("），2K–12K 加速 ", False),
    ("1.59–5.37×", True),
    ("，3601/3601 回归及 1000 次 bitwise 确定性通过，短序列无显著退化并交付 wheel。", False),
]


INFERENCE_PARTS = [
    ("推理算子支持：", True),
    (
        "针对 batch-invariant MoE Router 在小 M decode 中的无效计算及 CUDA Graph 动态选核风险，在 "
        "vLLM 中工程化接入 DeepGEMM/Full-K/persistent 三级后端，负责 tensor-signature selector、capture "
        "前数值/Graph preflight、cache/fallback 及 workspace 生命周期；48 层 Router GEMM 中位耗时下降 ",
        False,
    ),
    ("73.39%", True),
    ("、每 step 的 device-kernel 时长之和下降 ", False),
    ("6.16%", True),
    ("，135/135 kernel 与 20/20 模型场景 bitwise 一致；prefix-cache hit 下 TP1/TP2 整请求处理性能提升 ", False),
    ("3.81%/4.36%", True),
    ("，与异步解码联合提升 ", False),
    ("11.05%/8.26%", True),
    ("。", False),
]


def iter_unique_cells(document):
    seen = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                key = cell._tc
                if key in seen:
                    continue
                seen.add(key)
                yield cell


def iter_paragraphs(document):
    yield from document.paragraphs
    for cell in iter_unique_cells(document):
        yield from cell.paragraphs
        for nested in cell.tables:
            for row in nested.rows:
                for nested_cell in row.cells:
                    yield from nested_cell.paragraphs


def clear_paragraph(paragraph):
    p = paragraph._p
    p_pr = p.pPr
    saved = deepcopy(p_pr) if p_pr is not None else None
    for child in list(p):
        p.remove(child)
    if saved is not None:
        p.insert(0, saved)


def set_parts(paragraph, parts):
    clear_paragraph(paragraph)
    for text, bold in parts:
        run = paragraph.add_run(text)
        run.bold = bold


def remove_paragraph(paragraph):
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def find_one(document, prefix):
    matches = [p for p in iter_paragraphs(document) if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def split_table_with_page_break(table, row_index):
    """Split a table before row_index and insert a true page break between halves."""
    source = table._tbl
    cloned = deepcopy(source)
    source_rows = list(source.findall(qn("w:tr")))
    cloned_rows = list(cloned.findall(qn("w:tr")))
    if not 0 < row_index < len(source_rows):
        raise ValueError(f"invalid split row {row_index} for {len(source_rows)} rows")

    for row in source_rows[row_index:]:
        source.remove(row)
    for row in cloned_rows[:row_index]:
        cloned.remove(row)

    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    run.append(page_break)
    paragraph.append(run)

    source.addnext(paragraph)
    paragraph.addnext(cloned)


def main():
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)

    document = Document(TEMPLATE)

    replacements = [
        ("FlashInfer CUDA Graph", RL_PARTS),
        ("R3 Router Replay", TRAINING_PARTS),
        ("H200 确定性 Router GEMM", INFERENCE_PARTS),
    ]
    for prefix, parts in replacements:
        paragraph = find_one(document, prefix)
        set_parts(paragraph, parts)

    old_oe = find_one(document, "OE 异步算子开发")
    remove_paragraph(old_oe)

    direction = find_one(document, "技术方向与语言：")
    set_parts(
        direction,
        [
            ("技术方向与语言：", True),
            ("聚焦 RL Infra、GPU 算子与大模型推理系统；IELTS 6.5、CET-6，持有华为 HCIA-AI 认证。", False),
        ],
    )

    # The source keeps internships, open-source work and projects in one table.
    # Split it before the projects section so its heading cannot become an orphan.
    split_table_with_page_break(document.tables[1], 6)

    # Avoid stale pagination hints from the source Word rendering; pagination is
    # verified again after the new document is converted to PDF.
    for rendered_break in document._element.xpath(".//w:lastRenderedPageBreak"):
        parent = rendered_break.getparent()
        if parent is not None:
            parent.remove(rendered_break)

    document.core_properties.title = "冯浩然 - AI Infra 中文简历（RL/训练/推理算子版）"
    document.core_properties.subject = "两页中文简历确认稿"
    document.core_properties.author = "冯浩然"
    document.core_properties.last_modified_by = "冯浩然"
    document.core_properties.comments = "基于 2026-09-11 开发记录增量重写；原始模板未覆盖。"
    document.core_properties.modified = datetime.now(timezone.utc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
