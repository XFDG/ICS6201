"""Build matching Word/PDF resume from the flat output Markdown."""
from pathlib import Path
import re
import subprocess
import os
import argparse
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser()
parser.add_argument('source', type=Path)
parser.add_argument('--intro', action='store_true')
args = parser.parse_args()
SOURCE = args.source.resolve()
OUT = SOURCE.with_suffix('.docx')
ENGLISH = SOURCE.name.startswith('Haoran_')
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(1.1), Cm(1.2)
s.left_margin = s.right_margin = Cm(1.1)
style = d.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(9)
style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')

def font(run, size=9, bold=False, heading=False):
    run.font.name = 'Arial' if heading and ENGLISH else '黑体' if heading else 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '黑体' if heading else '宋体')

def inline(p, text, size=9, bold=False, heading=False):
    for token in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\))', text):
        if not token:
            continue
        if token.startswith('**'):
            inline(p, token[2:-2], size, True, heading)
        elif token.startswith('[') and '](' in token:
            label, url = token[1:-1].split('](', 1)
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
            r = p.add_run(label)
            font(r, size, bold, heading)
            color = OxmlElement('w:color'); color.set(qn('w:val'),'0563C1'); r._element.rPr.append(color)
            link.append(r._element); p._p.append(link)
        else:
            font(p.add_run(token), size, bold, heading)

def paragraph():
    p=d.add_paragraph()
    f=p.paragraph_format
    f.space_before=Pt(0); f.space_after=Pt(8 if args.intro else 4)
    f.line_spacing=1.24
    e=OxmlElement('w:snapToGrid'); e.set(qn('w:val'),'0'); p._p.get_or_add_pPr().append(e)
    return p

for block in SOURCE.read_text().split('\n\n'):
    block=block.strip()
    if not block: continue
    if block=='<!-- pagebreak -->':
        d.add_page_break(); continue
    for text in block.splitlines():
        p=paragraph()
        if text.startswith('# '):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,text[2:],16 if args.intro else 26,True)
        elif text.startswith('## '):
            p.paragraph_format.space_before=Pt(6)
            p.paragraph_format.keep_with_next=True
            inline(p,text[3:],10,True,True)
            border=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
            for key,val in [('val','single'),('sz','4'),('color','888888')]: b.set(qn('w:'+key),val)
            border.append(b); p._p.get_or_add_pPr().append(border)
        elif text.startswith('### '):
            label, date=text[4:].rsplit(' | ',1)
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(18.8), WD_TAB_ALIGNMENT.RIGHT)
            logo='iquest.png' if label.startswith(('九坤','Ubiquant')) else 'moorethreads.png' if label.startswith(('摩尔','Moore Threads')) else None
            if logo:
                p.add_run().add_picture(str(ROOT/'doc/assets/resume_20260929'/logo),height=Pt(12))
                p.add_run('  ')
            inline(p,label,9,True,True)
            inline(p,'\t'+date,8.5)
        else:
            if text.startswith('- '):
                text='• '+text[2:]
                p.paragraph_format.left_indent=Pt(8)
                p.paragraph_format.first_line_indent=Pt(-8)
            if text.startswith(('电话：','Phone:')): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,text,11 if args.intro else 9)
            if text.startswith('**') and text.endswith('**'): p.paragraph_format.keep_with_next=True
d.core_properties.title=SOURCE.stem
d.core_properties.author='冯浩然'
d.save(OUT)
profile=ROOT/'tmp/pdfs/release_word_profile'
subprocess.run(['libreoffice',f'-env:UserInstallation={profile.as_uri()}','--headless','--convert-to','pdf','--outdir',str(OUT.parent),str(OUT)],env={**os.environ,'SAL_USE_VCLPLUGIN':'svp'},check=True)
# Some LibreOffice profiles omit hyperlink annotations; restore them from labels.
import pdfplumber
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Link
pdf = OUT.with_suffix('.pdf')
writer = PdfWriter(clone_from=PdfReader(pdf))
used = set()
with pdfplumber.open(pdf) as pages:
    for label, url in re.findall(r'\[([^\]]+)\]\(([^)]+)\)', SOURCE.read_text()):
        found = False
        pattern = r'\s*'.join(re.escape(c) for c in label if not c.isspace())
        for idx, page in enumerate(pages.pages):
            for hit in page.search(pattern):
                key=(idx,round(hit['x0'],1),round(hit['top'],1))
                if key in used: continue
                used.add(key)
                writer.add_annotation(idx,Link(rect=(hit['x0'],page.height-hit['bottom'],hit['x1'],page.height-hit['top']),url=url))
                found=True
                break
            if found: break
        if not found: raise RuntimeError(f'Link label not found: {label}')
with pdf.open('wb') as target: writer.write(target)
print(OUT)
